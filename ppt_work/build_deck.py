"""
Builds the SIH 2026 idea deck (PS 26171) on the official template.
Keeps the template's slide order, titles and pointer headings, fills each with
diagram/card layouts, and drops the instructions slide (6 slide limit).
Run: python build_deck.py
"""
import copy

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC = "../SIH2026-IDEA-Presentation-Format.pptx"
OUT = "../SIH2026-IDEA-Presentation-HexaBits.pptx"
ICONS = "icons"

TEAM = "HexaBits"
IDEA = "Fusion Privacy Agent"

# palette pulled from the SIH logo + template footer blue
BLUE = RGBColor(0x00, 0x70, 0xC0)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
SAFFRON = RGBColor(0xE8, 0x77, 0x22)
GREEN = RGBColor(0x1E, 0x8E, 0x4A)
INK = RGBColor(0x26, 0x2E, 0x3B)
MUTED = RGBColor(0x5B, 0x66, 0x75)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE_TINT = RGBColor(0xEA, 0xF3, 0xFB)
SAFFRON_TINT = RGBColor(0xFD, 0xF0, 0xE6)
GREEN_TINT = RGBColor(0xE7, 0xF5, 0xEC)
GREY_TINT = RGBColor(0xF3, 0xF5, 0xF8)
LINE = RGBColor(0xC9, 0xD3, 0xDF)

prs = Presentation(SRC)


# ---------- helpers ----------

def find(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    return None


def drop(shape):
    shape._element.getparent().remove(shape._element)


def rect(slide, x, y, w, h, fill=None, line=None, radius=None, dash=False, line_w=1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = 4  # dash
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def _style_run(r, size, bold, color, italic, font):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font


def _bullet(p, color, indent=0.16):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buClr", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", str(color))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def text(slide, x, y, w, h, paras, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, font="Calibri", space_after=0, bullet_color=None,
         shape=None):
    """paras: list of str | list of (text, {overrides}) runs. Prefix a str/run-list with
    dict {'bullet': True} as first element to bullet that paragraph."""
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = para
        is_bullet = False
        if isinstance(para, tuple) and para and isinstance(para[0], dict):
            is_bullet = para[0].get("bullet", False)
            runs = para[1]
        if isinstance(runs, str):
            runs = [(runs, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            _style_run(r, o.get("size", size), o.get("bold", bold), o.get("color", color),
                       o.get("italic", italic), o.get("font", font))
        if is_bullet:
            _bullet(p, bullet_color or color)
    return shape


def B(runs):
    return ({"bullet": True}, runs)


def icon_disc(slide, x, y, d, fill, icon):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    pad = d * 0.24
    slide.shapes.add_picture(f"{ICONS}/{icon}.png", Inches(x + pad), Inches(y + pad),
                             Inches(d - 2 * pad), Inches(d - 2 * pad))


def arrow(slide, x1, y1, x2, y2, color=MUTED, w=1.5, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = 4
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def section(slide, x, y, w, label, color=BLUE):
    """Template pointer heading, kept verbatim as the section label."""
    text(slide, x, y, w, 0.3, [[(label, {})]], size=13, bold=True, color=color)


def header(slide, title, tagline, title_shape="Title 1", oval_names=("Oval 9", "Oval 10", "Oval 11", "Oval 8")):
    t = find(slide, title_shape)
    t.left, t.top, t.width, t.height = Inches(1.9), Inches(0.12), Inches(8.6), Inches(0.62)
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = tf.margin_bottom = 0
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    for br in p._p.findall(qn("a:br")):
        p._p.remove(br)
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].text = title
    p.runs[0].font.size = Pt(30)
    for r in p.runs[1:]:
        r.text = ""
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    text(slide, 1.9, 0.76, 8.6, 0.32, tagline, size=13, italic=True, color=MUTED, align=PP_ALIGN.CENTER)
    for n in oval_names:
        o = find(slide, n)
        if o is not None:
            tf = o.text_frame
            tf.word_wrap = False
            tf.margin_left = tf.margin_right = 0
            p = tf.paragraphs[0]
            p.runs[0].text = TEAM
            p.runs[0].font.size = Pt(15)
            p.runs[0].font.bold = True
            for r in p.runs[1:]:
                r.text = ""
            for extra in tf.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)


S = list(prs.slides)

# ---------- slide 1: title ----------
s1 = S[0]
tb = find(s1, "TextBox 9")
tb.top = Inches(2.3)
tb.height = Inches(4.4)
tf = tb.text_frame
src_p = next(p for p in tf.paragraphs if p.runs)
tmpl_pPr = copy.deepcopy(src_p._p.find(qn("a:pPr")))
tmpl_rPr = copy.deepcopy(src_p.runs[0]._r.find(qn("a:rPr")))
for p in list(tf.paragraphs)[1:]:
    p._p.getparent().remove(p._p)
first = tf.paragraphs[0]
for r in list(first.runs):
    r._r.getparent().remove(r._r)

fields = [
    ("Problem Statement ID – ", "SIH26171"),
    ("Problem Statement Title – ", "On-device Visual Perception for Light-weight Browser Agents"),
    ("Theme – ", "Miscellaneous"),
    ("PS Category – ", "Software"),
    ("Team ID – ", ""),
    ("Team Name – ", TEAM),
]
for i, (label, val) in enumerate(fields):
    p = first if i == 0 else tf.add_paragraph()
    if tmpl_pPr is not None:
        old = p._p.find(qn("a:pPr"))
        if old is not None:
            p._p.remove(old)
        p._p.insert(0, copy.deepcopy(tmpl_pPr))
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.05
    p.space_after = Pt(10)
    for t, b in ((label, True), (val, False)):
        if not t:
            continue
        r = p.add_run()
        if tmpl_rPr is not None:
            r._r.insert(0, copy.deepcopy(tmpl_rPr))
        r.text = t
        r.font.size = Pt(19)
        r.font.bold = True if b else False
        r.font.color.rgb = NAVY if b else INK

# ---------- slide 2: idea ----------
s2 = S[1]
header(s2, IDEA.upper(), "See locally. Redact locally. Only then ask the server.")
drop(find(s2, "TextBox 8"))

section(s2, 0.45, 1.22, 8, "Proposed Solution (Describe your Idea/Solution/Prototype)")
text(s2, 0.45, 1.52, 12.4, 0.3,
     "A browser extension that reads the screen with on-device vision models, masks every sensitive region "
     "locally, and lets a server VLM plan actions from the masked view only.",
     size=12, color=MUTED)

steps = [
    ("eye", BLUE, "SEE", "Florence-2 (ViT), YOLOv8n and OCR read the screen on WebGPU"),
    ("card", SAFFRON, "DETECT", "DOM tags, regex, NER and BlazeFace find PII and faces"),
    ("eyeSlash", SAFFRON, "REDACT", "Masks drawn on-device before any network request"),
    ("server", NAVY, "PLAN", "Server VLM gets masked image + redaction manifest"),
    ("click", GREEN, "ACT", "Extension clicks, types, scrolls, loops until done"),
]
bx, bw, gap, by, bh = 0.45, 2.24, 0.30, 2.02, 1.38
for i, (ic, col, head, body) in enumerate(steps):
    x = bx + i * (bw + gap)
    tint = {BLUE: BLUE_TINT, SAFFRON: SAFFRON_TINT, NAVY: GREY_TINT, GREEN: GREEN_TINT}[col]
    rect(s2, x, by, bw, bh, fill=tint, radius=0.1)
    icon_disc(s2, x + 0.14, by + 0.14, 0.46, col, ic)
    text(s2, x + 0.7, by + 0.2, bw - 0.8, 0.34, [[(f"{i + 1}  {head}", {})]], size=13, bold=True, color=col)
    text(s2, x + 0.14, by + 0.68, bw - 0.28, 0.66, body, size=10.5, color=INK)
    if i < len(steps) - 1:
        arrow(s2, x + bw + 0.03, by + bh / 2, x + bw + gap - 0.03, by + bh / 2, color=MUTED)
# device boundary between REDACT and PLAN
bnd_x = bx + 3 * (bw + gap) - gap / 2
s = rect(s2, bnd_x - 0.01, by - 0.12, 0.02, bh + 0.24, fill=SAFFRON)
text(s2, bnd_x - 1.6, by + bh + 0.08, 3.2, 0.24, "device boundary: only masked data crosses",
     size=9.5, italic=True, color=SAFFRON, align=PP_ALIGN.CENTER)

# lower row
ly = 3.82
section(s2, 0.45, ly, 5.2, "How it addresses the problem")
addr = [
    [("Sensitive pixels never leave the browser: ", {"bold": True}), ("passwords, IDs, faces and PII are masked on-device first.", {})],
    [("Server stays useful: ", {"bold": True}), ("the redaction manifest tells it what kind of field sits where, so it can still plan.", {})],
    [("Light enough for a laptop: ", {"bold": True}), ("quantized models, WebGPU with WASM fallback, all bundled offline.", {})],
    [("End-to-end: ", {"bold": True}), ("a multi-step agent loop completes real form tasks, with user confirmation on submit or pay.", {})],
]
text(s2, 0.45, ly + 0.38, 5.3, 2.6, [B(r) for r in addr], size=11.5, space_after=7, bullet_color=BLUE)

section(s2, 6.1, ly, 6.5, "Innovation and uniqueness of the solution")
inno = [
    ("layers", BLUE, "Confidence fusion", "Detectors vote with noisy-OR; a weak lone hit is still masked (fail-safe)."),
    ("check", GREEN, "DOM-wins safe zones", "DOM structure overrides vision false positives, cutting over-redaction."),
    ("shield", NAVY, "Redaction-aware server", "Guard blocks the planner from typing into fields it cannot see."),
    ("vault", SAFFRON, "Local answer vault", "Values the agent asks for are filled on-device, never sent to server."),
]
cw, ch = 3.28, 1.24
for i, (ic, col, head, body) in enumerate(inno):
    cx = 6.1 + (i % 2) * (cw + 0.16)
    cy = ly + 0.4 + (i // 2) * (ch + 0.14)
    rect(s2, cx, cy, cw, ch, fill=GREY_TINT, radius=0.1)
    icon_disc(s2, cx + 0.12, cy + 0.14, 0.42, col, ic)
    text(s2, cx + 0.64, cy + 0.2, cw - 0.72, 0.3, head, size=12, bold=True, color=NAVY)
    text(s2, cx + 0.12, cy + 0.64, cw - 0.24, 0.58, body, size=10.5, color=INK)

# ---------- slide 3: technical approach ----------
s3 = S[2]
header(s3, "TECHNICAL APPROACH", "Heavy perception on the client, reasoning on the server, raw data never crosses.")
drop(find(s3, "TextBox 8"))

section(s3, 0.45, 1.22, 4.3, "Technologies to be used")
groups = [
    ("On-device vision / NLP", SAFFRON, ["Florence-2 (ViT)", "YOLOv8n", "BlazeFace", "BERT-NER", "Tesseract OCR", "Whisper-tiny"]),
    ("In-browser runtime", BLUE, ["WebGPU", "WASM fallback", "ONNX Runtime Web", "Transformers.js", "MediaPipe"]),
    ("Extension", NAVY, ["TypeScript", "Manifest V3", "Offscreen document", "Side panel UI", "Vite"]),
    ("Server", GREEN, ["Python", "FastAPI", "Open-weights VLM", "OpenAI-compatible API", "Pydantic"]),
]
gy = 1.6
for title_, col, chips in groups:
    text(s3, 0.45, gy, 4.2, 0.26, title_, size=11, bold=True, color=col)
    gy += 0.3
    cx, row_y = 0.45, gy
    for chip in chips:
        cwid = 0.13 + len(chip) * 0.078
        if cx + cwid > 4.65:
            cx = 0.45
            row_y += 0.36
        c = rect(s3, cx, row_y, cwid, 0.28, fill=WHITE, line=col, radius=0.5, line_w=1)
        text(s3, 0, 0, 0, 0, chip, size=9.5, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=c)
        cx += cwid + 0.08
    gy = row_y + 0.46

section(s3, 5.0, 1.22, 7.9, "Methodology and process for implementation")

# browser zone
zx, zw = 5.0, 7.88
rect(s3, zx, 1.58, zw, 3.2, fill=SAFFRON_TINT, radius=0.04)
text(s3, zx + 0.15, 1.64, 5, 0.26, "USER'S BROWSER  (trusted, on-device)", size=10, bold=True, color=SAFFRON)

col_x = [zx + 0.2, zx + 2.84, zx + 5.48]
box_w = 2.2


def node(slide, x, y, w, h, head, body, col, fill=WHITE):
    rect(slide, x, y, w, h, fill=fill, line=col, radius=0.1, line_w=1.25)
    text(slide, x + 0.1, y + 0.07, w - 0.2, 0.26, head, size=11, bold=True, color=col)
    text(slide, x + 0.1, y + 0.36, w - 0.2, h - 0.4, body, size=9.5, color=INK)


r1y, r2y, nh = 1.98, 3.4, 1.0
node(s3, col_x[0], r1y, box_w, nh, "1. Capture", "Screenshot + DOM skeleton (tags, types, labels; no values)", NAVY)
node(s3, col_x[1], r1y, box_w, nh, "2. Perceive", "Offscreen doc: Florence-2, YOLO, OCR, NER, BlazeFace on WebGPU", BLUE)
node(s3, col_x[2], r1y, box_w, nh, "3. Fuse + Redact", "DOM + vision boxes, noisy-OR scoring, black-box masks", SAFFRON)
arrow(s3, col_x[0] + box_w, r1y + nh / 2, col_x[1], r1y + nh / 2)
arrow(s3, col_x[1] + box_w, r1y + nh / 2, col_x[2], r1y + nh / 2)

node(s3, col_x[0], r2y, box_w, nh, "6. Risk gate + vault", "Submit / pay / delete wait for user OK; asked values filled locally", SAFFRON)
node(s3, col_x[1], r2y, box_w, nh, "7. Execute + loop", "Click, type, scroll, navigate; observe again until done (max 15)", GREEN)
arrow(s3, col_x[0] + box_w, r2y + nh / 2, col_x[1], r2y + nh / 2)

# privacy boundary
by_ = 4.97
rect(s3, zx, by_ - 0.01, zw, 0.02, fill=SAFFRON)
text(s3, zx + 0.1, by_ + 0.04, 2.6, 0.22, "privacy boundary", size=9, italic=True, color=SAFFRON)

# server zone
rect(s3, zx, 5.2, zw, 1.62, fill=GREY_TINT, radius=0.06)
text(s3, col_x[1], 5.24, box_w + 0.5, 0.26, "SERVER  (sees only sanitized context)", size=10, bold=True, color=NAVY)
sy, sh = 5.54, 1.08
node(s3, col_x[2], sy, box_w, sh, "4. /next-action", "FastAPI receives masked PNG + manifest + DOM skeleton", NAVY)
node(s3, col_x[1], sy, box_w, sh, "5a. VLM planner", "Open-weights VLM, prompt knows the redaction scheme", NAVY)
node(s3, col_x[0], sy, box_w, sh, "5b. Safety guard", "Rejects typing into masked fields and made-up selectors", NAVY)
arrow(s3, col_x[2], sy + sh / 2, col_x[1] + box_w, sy + sh / 2)
arrow(s3, col_x[1], sy + sh / 2, col_x[0] + box_w, sy + sh / 2)

# crossing arrows
dx = col_x[2] + box_w / 2
arrow(s3, dx, r1y + nh, dx, sy, color=SAFFRON, w=2)
text(s3, dx + 0.1, r2y + 0.05, 1.05, 1.05,
     [[("sanitized", {"bold": True})], "masked image", "manifest", "DOM skeleton"],
     size=9.5, color=SAFFRON)
ux = col_x[0] + box_w / 2
arrow(s3, ux, sy, ux, r2y + nh, color=GREEN, w=2)
text(s3, ux + 0.1, 4.6, 1.2, 0.26, [[("action JSON", {"bold": True})]], size=9.5, color=GREEN)

# ---------- slide 4: feasibility ----------
s4 = S[3]
header(s4, "FEASIBILITY AND VIABILITY", "Designed around the five evaluation metrics, working prototype already running.")
drop(find(s4, "TextBox 8"))

section(s4, 0.45, 1.22, 6.6, "Analysis of the feasibility of the idea")
metrics = [
    ("25%", "Visual context accuracy", "Florence-2 captions + YOLO elements fused with the real DOM list, so actions hit real selectors"),
    ("20%", "PII recall / precision", "DOM + regex (Luhn, Verhoeff) + NER + BlazeFace fused; eval: 100% recall, 91.3% precision"),
    ("20%", "Redaction precision", "Pixel-exact DOM boxes; eval: 92.8% of masked pixels are sensitive, mean IoU 0.99"),
    ("20%", "Client resources", "int8 / fp16 models in one offscreen doc; heap and WebGPU shown live per step"),
    ("15%", "End-to-end latency", "Per-step breakdown shown live: capture, on-device, planner, act"),
]
my = 1.6
for w_, name, how in metrics:
    rect(s4, 0.45, my, 6.6, 0.9, fill=GREY_TINT, radius=0.12)
    d = rect(s4, 0.58, my + 0.17, 0.72, 0.56, fill=BLUE, radius=0.3)
    text(s4, 0, 0, 0, 0, w_, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=d)
    text(s4, 1.46, my + 0.12, 5.45, 0.26, name, size=12.5, bold=True, color=NAVY)
    text(s4, 1.46, my + 0.41, 5.45, 0.46, how, size=10.5, color=INK)
    my += 1.02

section(s4, 7.35, 1.22, 5.5, "Potential challenges and risks", color=SAFFRON)
text(s4, 10.2, 1.22, 2.7, 0.3, [[("Strategies for overcoming", {})]], size=13, bold=True, color=GREEN)
risks = [
    ("Missed PII (false negative)", "Fail-safe: a low-confidence lone hit is still masked"),
    ("Over-redaction hides context", "Manifest keeps field type; DOM-safe zones"),
    ("Heavy models, low-end laptops", "Quantized, lazy load, WASM fallback"),
    ("Planner hallucinates actions", "Selectors only from real DOM; server guard"),
    ("Site CSP blocks ML workers", "ML runs in extension's offscreen document"),
]
ry = 1.6
for risk, fix in risks:
    rect(s4, 7.35, ry, 2.62, 0.9, fill=SAFFRON_TINT, radius=0.12)
    icon_disc(s4, 7.45, ry + 0.25, 0.4, SAFFRON, "warn")
    text(s4, 7.95, ry + 0.08, 1.95, 0.74, risk, size=11, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s4, 9.99, ry + 0.45, 10.19, ry + 0.45, color=MUTED)
    rect(s4, 10.2, ry, 2.68, 0.9, fill=GREEN_TINT, radius=0.12)
    icon_disc(s4, 10.3, ry + 0.25, 0.4, GREEN, "check")
    text(s4, 10.8, ry + 0.08, 2.0, 0.74, fix, size=11, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    ry += 1.02

# ---------- slide 5: impact ----------
s5 = S[4]
header(s5, "IMPACT AND BENEFITS", "Makes AI agents usable exactly where privacy matters most.")
drop(find(s5, "TextBox 8"))

stats = [("0", 0.6, "raw field values or PII pixels sent to the server"),
         ("6", 0.6, "on-device models bundled in the extension, no CDN at runtime"),
         ("< 1 ms", 1.75, "to fuse all detections into the redaction plan")]
sx = 0.45
for big, nw, lbl in stats:
    rect(s5, sx, 1.25, 4.0, 1.1, fill=NAVY, radius=0.1)
    text(s5, sx + 0.25, 1.3, nw, 1.0, big, size=34, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE,
         font="Cambria")
    text(s5, sx + 0.35 + nw, 1.3, 3.5 - nw, 1.0, lbl, size=12, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    sx += 4.2

section(s5, 0.45, 2.6, 6.2, "Potential impact on the target audience")
aud = [
    ("landmark", BLUE, "Citizens on e-governance portals",
     "Agent helps fill Aadhaar, PAN, bank and scholarship forms while ID numbers stay on the device."),
    ("shield", NAVY, "Government and enterprise staff",
     "Cloud AI on internal portals without exposing confidential screens or employee data."),
    ("access", GREEN, "Low-literacy and differently-abled users",
     "On-device voice input (Whisper) and a side panel make web tasks hands-free."),
]
ay = 3.0
for ic, col, head, body in aud:
    icon_disc(s5, 0.45, ay + 0.05, 0.56, col, ic)
    text(s5, 1.15, ay, 5.4, 0.3, head, size=12.5, bold=True, color=NAVY)
    text(s5, 1.15, ay + 0.32, 5.4, 0.7, body, size=11.5, color=INK)
    ay += 1.22

section(s5, 6.95, 2.6, 5.9, "Benefits of the solution (social, economic, environmental, etc.)")
ben = [
    ("users", BLUE, "Social", "Trust in AI help; data minimisation in line with the DPDP Act 2023."),
    ("rupee", GREEN, "Economic", "Perception runs on the client; open-weights server model, no vendor lock-in."),
    ("lock", SAFFRON, "Security", "Server never stores raw PII, so a server breach leaks nothing personal."),
    ("bolt", NAVY, "Environmental", "Less server GPU work and smaller payloads per step."),
]
bw2, bh2 = 2.88, 1.62
for i, (ic, col, head, body) in enumerate(ben):
    cx = 6.95 + (i % 2) * (bw2 + 0.14)
    cy = 3.0 + (i // 2) * (bh2 + 0.14)
    rect(s5, cx, cy, bw2, bh2, fill=GREY_TINT, radius=0.08)
    icon_disc(s5, cx + 0.14, cy + 0.14, 0.44, col, ic)
    text(s5, cx + 0.68, cy + 0.2, bw2 - 0.8, 0.32, head, size=13, bold=True, color=col)
    text(s5, cx + 0.14, cy + 0.7, bw2 - 0.28, 0.9, body, size=11.5, color=INK)

# ---------- slide 6: references ----------
s6 = S[5]
header(s6, "RESEARCH AND REFERENCES", "Every model and runtime below is open-source or open-weights.")
drop(find(s6, "TextBox 8"))
section(s6, 0.45, 1.22, 8, "Details / Links of the reference and research work")

refcols = [
    ("On-device models", SAFFRON, [
        ("Florence-2 (ViT)", "Xiao et al., CVPR 2024, arXiv:2311.06242"),
        ("BlazeFace", "Bazarevsky et al., arXiv:1907.05047"),
        ("YOLOv8", "github.com/ultralytics/ultralytics"),
        ("BERT NER", "Devlin et al., arXiv:1810.04805; hf.co/Xenova/bert-base-NER"),
        ("Whisper", "Radford et al., arXiv:2212.04356"),
        ("Tesseract OCR", "github.com/naptha/tesseract.js"),
    ]),
    ("Runtimes and platform", BLUE, [
        ("ONNX Runtime Web", "onnxruntime.ai/docs/tutorials/web"),
        ("Transformers.js", "huggingface.co/docs/transformers.js"),
        ("WebGPU", "w3.org/TR/webgpu"),
        ("MediaPipe Tasks", "ai.google.dev/edge/mediapipe"),
        ("Chrome Offscreen API", "developer.chrome.com/docs/extensions"),
        ("FastAPI", "fastapi.tiangolo.com"),
    ]),
    ("Web agents and policy", NAVY, [
        ("SeeAct", "Zheng et al., ICML 2024, arXiv:2401.01614"),
        ("Mind2Web", "Deng et al., NeurIPS 2023, arXiv:2306.06070"),
        ("Llama 4 Scout (open-weights VLM)", "hf.co/meta-llama/Llama-4-Scout-17B-16E-Instruct"),
        ("DPDP Act 2023", "meity.gov.in, Digital Personal Data Protection Act"),
        ("SIH 2026 PS 26171", "ISRO / SAC, On-device Visual Perception for Browser Agents"),
    ]),
]
colw = 4.0
for i, (head, col, items) in enumerate(refcols):
    x = 0.45 + i * (colw + 0.22)
    rect(s6, x, 1.62, colw, 3.85, fill=GREY_TINT, radius=0.04)
    icon_disc(s6, x + 0.16, 1.76, 0.42, col, ["cpu", "code", "book"][i])
    text(s6, x + 0.7, 1.8, colw - 0.8, 0.34, head, size=13, bold=True, color=col)
    paras = []
    for name, ref in items:
        paras.append([(name, {"bold": True, "color": NAVY, "size": 11.5})])
        paras.append([(ref, {"color": MUTED, "size": 10})])
    text(s6, x + 0.18, 2.3, colw - 0.32, 3.1, paras, size=10, space_after=2)

st = rect(s6, 0.45, 5.62, 12.44, 1.08, fill=NAVY, radius=0.1)
icon_disc(s6, 0.65, 5.85, 0.62, GREEN, "flask")
text(s6, 1.45, 5.72, 11.2, 0.3, "Our own research: working prototype + reproducible eval harness", size=13, bold=True, color=WHITE)
text(s6, 1.45, 6.04, 11.2, 0.6,
     "Extension, server and Jarvis desktop companion built. Eval (eval/results.md) runs the shipping redaction code on "
     "32 labelled items: 100% recall, 91.3% precision, 92.8% pixel precision; text PII 100% / 100% on 25 cases.",
     size=11, color=WHITE)


# ---------- drop instructions slide (template allows max 6) ----------
sldIdLst = prs.slides._sldIdLst
last = list(sldIdLst)[6]
prs.part.drop_rel(last.get(qn("r:id")))
sldIdLst.remove(last)

prs.save(OUT)
print("saved", OUT)
